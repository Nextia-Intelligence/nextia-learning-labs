"""Spreadsheets: read the cells themselves, with their units and their addresses.

    python -m doc_extract.sheets D057            # the supplier's price list as a spreadsheet
    python -m doc_extract.sheets D057 --naive    # the cells as they are, without the merged cells

Never turn a spreadsheet into a picture and read it with OCR: the cells already hold exact values.
What needs care: the header is not on row 1 (a title is above it), the units are in the header
("Price each (CAD)", "Pack weight (kg)"), and the category is one merged cell over several rows.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass
class SheetCell:
    value: object
    sheet: str
    cell: str  # the address, e.g. "E12"
    merged: str = ""  # the merged range the value comes from, e.g. "A5:A8"


def header_row(ws) -> int:
    """The first row with at least 3 text cells followed by a row with a number: the header."""
    for r in range(1, min(ws.max_row, 20) + 1):
        texts = [c.value for c in ws[r] if isinstance(c.value, str) and c.value.strip()]
        below = [c.value for c in ws[r + 1] if isinstance(c.value, (int, float))] if r < ws.max_row else []
        if len(texts) >= 3 and below:
            return r
    return 1


def unit_of(name: str) -> str | None:
    m = re.search(r"\(([^)]+)\)\s*$", name)
    return m[1] if m else None


def read(path: Path, sheet: str | None = None, fill_merged: bool = True) -> tuple[list[str], dict[str, str], list[dict[str, SheetCell]]]:
    """(column names, units, records). A record maps each column to its cell, with the address."""
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    merged: dict[str, tuple[object, str]] = {}
    if fill_merged:
        for rng in ws.merged_cells.ranges:
            top = ws.cell(rng.min_row, rng.min_col).value
            for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
                for c in row:
                    merged[c.coordinate] = (top, str(rng))
    h = header_row(ws)
    names = [str(c.value).strip() if c.value is not None else f"column {get_column_letter(c.column)}" for c in ws[h]]
    units = {n: u for n in names if (u := unit_of(n))}
    records = []
    for r in range(h + 1, ws.max_row + 1):
        cells = list(ws[r])
        if all(c.value is None for c in cells):
            break  # a blank row ends the table (notes may follow)
        rec = {}
        for name, c in zip(names, cells):
            value, rng = c.value, ""
            if value is None and c.coordinate in merged:
                value, rng = merged[c.coordinate]
            rec[name] = SheetCell(None if value == "n/a" else value, ws.title, c.coordinate, rng)
        records.append(rec)
    wb.close()
    return names, units, records


def main(argv: list[str] | None = None) -> None:
    from doc_extract import documents

    argv = sys.argv[1:] if argv is None else argv
    d = documents.get(argv[0] if argv else "D057")
    naive = "--naive" in argv
    names, units, records = read(d.path, fill_merged=not naive)
    first = records[0][names[0]].cell
    print(f"{d.file}: sheet {records[0][names[0]].sheet}, header on row {int(re.sub(r'[A-Z]', '', first)) - 1}")
    print(f"columns: {names}")
    print(f"units from the header: {units}")
    for rec in records[:3] + records[4:6]:
        print("  " + ", ".join(f"{k.split(' (')[0]}={c.value} [{c.sheet}!{c.cell}{' from ' + c.merged if c.merged else ''}]"
                                for k, c in list(rec.items())[:2] + list(rec.items())[4:5]))
    truth = d.truth.get("rows", [])
    ok = total = 0
    by_sku = {r["SKU"].value: r for r in records}
    for t in truth:
        got = by_sku.get(t["sku"], {})
        for key, col in (("category", "Category"), ("price_each_cad", "Price each (CAD)"),
                         ("price_case_cad", "Price per case (CAD)"), ("weight_kg", "Pack weight (kg)")):
            total += 1
            v = got[col].value if col in got else None
            want = t[key]
            if isinstance(v, (int, float)) and want is not None:
                ok += abs(float(v) - float(want)) < 0.005  # a number cell: compare the numbers
            else:
                ok += v == want
    print(f"{len(records)} products; cells correct: {ok} of {total} ({ok / total:.0%})")


if __name__ == "__main__":
    main()
